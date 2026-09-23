"""Local multimodal file indexing and safe review-queue sorting.

Scanning never moves, copies, deletes, or renames source files. Applying changes
requires records explicitly marked approved and verifies every copied file before
removing its source.
"""

from __future__ import annotations

import hashlib
import json
import mimetypes
import os
import shutil
import sqlite3
import subprocess
from collections.abc import Callable, Iterable
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

DOCUMENT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".json", ".csv", ".tsv", ".html", ".htm",
    ".xml", ".pdf", ".docx", ".doc", ".xlsx", ".pptx",
}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".heic"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".wmv", ".m4v", ".flv"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".flac", ".ogg", ".m4a", ".aac", ".wma", ".aiff", ".opus"}


class IndexerError(RuntimeError):
    """Base error for indexer failures."""


class UnsafePathError(IndexerError):
    """Raised when a source or destination escapes its configured root."""


@dataclass(frozen=True)
class IndexRecord:
    """One indexed file and its proposed destination."""

    id: int
    source_path: str
    target_path: str
    modality: str
    mime_type: str | None
    size_bytes: int
    modified_at: str
    sha256: str
    metadata: dict[str, Any]
    status: str
    error: str | None
    created_at: str
    updated_at: str

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


class LocalFileIndexer:
    """Index local media and manage a persistent human-review queue."""

    def __init__(
        self,
        database_path: Path,
        audit_logger: Any | None = None,
        analysis_hooks: Iterable[Callable[[Path, str, dict[str, Any]], dict[str, Any]]] | None = None,
    ) -> None:
        self.database_path = database_path
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self.audit_logger = audit_logger
        self.analysis_hooks = list(analysis_hooks or [])
        self._initialize()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        return connection

    def _initialize(self) -> None:
        with self._connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS indexed_files (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_path TEXT NOT NULL,
                    target_path TEXT NOT NULL,
                    modality TEXT NOT NULL,
                    mime_type TEXT,
                    size_bytes INTEGER NOT NULL,
                    modified_at TEXT NOT NULL,
                    sha256 TEXT NOT NULL,
                    metadata_json TEXT NOT NULL,
                    status TEXT NOT NULL,
                    error TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE(source_path, sha256)
                )
                """
            )
            connection.execute("CREATE INDEX IF NOT EXISTS idx_indexed_status ON indexed_files(status)")
            connection.execute("CREATE INDEX IF NOT EXISTS idx_indexed_hash ON indexed_files(sha256)")
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS index_roots (
                    id INTEGER PRIMARY KEY CHECK (id = 1),
                    source_root TEXT NOT NULL,
                    target_root TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )
                """
            )

    def scan(self, source_root: Path, target_root: Path) -> dict[str, Any]:
        """Scan files and populate the queue without filesystem mutations."""
        source = source_root.expanduser().resolve()
        target = target_root.expanduser().resolve()
        if not source.exists() or not source.is_dir():
            raise IndexerError(f"Source directory does not exist: {source}")
        if target == source or target.is_relative_to(source):
            raise UnsafePathError("Target directory must not be inside the source directory")
        with self._connect() as connection:
            connection.execute(
                "INSERT INTO index_roots(id, source_root, target_root, updated_at) VALUES (1, ?, ?, ?) "
                "ON CONFLICT(id) DO UPDATE SET source_root=excluded.source_root, target_root=excluded.target_root, updated_at=excluded.updated_at",
                (str(source), str(target), _now()),
            )
        records: list[IndexRecord] = []
        skipped = 0
        for path in self._iter_files(source):
            try:
                record = self._index_one(path, source, target)
                records.append(record)
            except (OSError, ValueError, IndexerError) as exc:
                skipped += 1
                self._record_error(path, source, target, exc)
        return {
            "source_root": str(source),
            "target_root": str(target),
            "scanned": len(records),
            "skipped": skipped,
            "pending": sum(record.status == "PENDING" for record in records),
            "duplicates": sum(record.status == "DUPLICATE" for record in records),
            "records": [record.to_dict() for record in records],
        }

    def _iter_files(self, root: Path) -> Iterable[Path]:
        for path in sorted(root.rglob("*")):
            if path.is_file() and not path.is_symlink() and self._classify(path) is not None:
                yield path

    def _index_one(self, path: Path, source_root: Path, target_root: Path) -> IndexRecord:
        source_path = self._safe_child(path, source_root)
        stat = source_path.stat()
        digest = sha256_file(source_path)
        modality = self._classify(source_path)
        if modality is None:
            raise IndexerError(f"Unsupported file type: {source_path.suffix}")
        metadata = extract_metadata(source_path, modality, self.analysis_hooks)
        target_path = self._propose_target(source_path, target_root, modality, digest)
        now = _now()
        with self._connect() as connection:
            connection.execute(
                "UPDATE indexed_files SET status='SUPERSEDED', updated_at=? "
                "WHERE source_path=? AND sha256<>? AND status IN ('PENDING', 'APPROVED')",
                (now, str(source_path), digest),
            )
        status = "DUPLICATE" if self._hash_exists_elsewhere(digest, str(source_path)) else "PENDING"
        with self._connect() as connection:
            connection.execute(
                """
                INSERT INTO indexed_files
                    (source_path, target_path, modality, mime_type, size_bytes, modified_at,
                     sha256, metadata_json, status, error, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, ?, ?)
                ON CONFLICT(source_path, sha256) DO UPDATE SET
                    target_path=excluded.target_path,
                    modality=excluded.modality,
                    mime_type=excluded.mime_type,
                    size_bytes=excluded.size_bytes,
                    modified_at=excluded.modified_at,
                    metadata_json=excluded.metadata_json,
                    status=CASE
                        WHEN indexed_files.status IN ('APPROVED', 'APPLIED', 'PARTIAL') THEN indexed_files.status
                        ELSE excluded.status
                    END,
                    error=NULL,
                    updated_at=excluded.updated_at
                """,
                (
                    str(source_path), str(target_path), modality, mimetypes.guess_type(source_path.name)[0],
                    stat.st_size, datetime.fromtimestamp(stat.st_mtime, UTC).isoformat(), digest,
                    json.dumps(metadata, ensure_ascii=False, sort_keys=True), status, now, now,
                ),
            )
            row = connection.execute(
                "SELECT * FROM indexed_files WHERE source_path = ? AND sha256 = ?",
                (str(source_path), digest),
            ).fetchone()
        assert row is not None
        record = _row_to_record(row)
        self._audit("index", record.to_dict())
        return record

    def _record_error(self, path: Path, source_root: Path, target_root: Path, error: Exception) -> None:
        now = _now()
        source = str(self._safe_child(path, source_root))
        target = str(target_root / path.name)
        with self._connect() as connection:
            connection.execute(
                """
                INSERT INTO indexed_files
                    (source_path, target_path, modality, mime_type, size_bytes, modified_at,
                     sha256, metadata_json, status, error, created_at, updated_at)
                VALUES (?, ?, 'unknown', NULL, 0, '', '', '{}', 'ERROR', ?, ?, ?)
                ON CONFLICT(source_path, sha256) DO UPDATE SET error=excluded.error, status='ERROR', updated_at=excluded.updated_at
                """,
                (source, target, str(error), now, now),
            )
        self._audit("index_error", {"source_path": source, "error": type(error).__name__})

    def list_queue(self, statuses: set[str] | None = None) -> list[IndexRecord]:
        """Return queue records, optionally filtered by status."""
        with self._connect() as connection:
            if statuses:
                values = sorted(statuses)
                placeholders = ",".join("?" for _ in values)
                rows = connection.execute(
                    f"SELECT * FROM indexed_files WHERE status IN ({placeholders}) ORDER BY id", values
                ).fetchall()
            else:
                rows = connection.execute("SELECT * FROM indexed_files ORDER BY id").fetchall()
        return [_row_to_record(row) for row in rows]

    def approve(self, record_ids: list[int] | None = None, approve_all: bool = False) -> int:
        """Mark pending records approved; this still performs no filesystem operation."""
        with self._connect() as connection:
            if approve_all:
                cursor = connection.execute(
                    "UPDATE indexed_files SET status='APPROVED', updated_at=? WHERE status='PENDING'", (_now(),)
                )
            else:
                if not record_ids:
                    raise IndexerError("Provide record IDs or approve_all=True")
                placeholders = ",".join("?" for _ in record_ids)
                cursor = connection.execute(
                    f"UPDATE indexed_files SET status='APPROVED', updated_at=? WHERE status='PENDING' AND id IN ({placeholders})",
                    [_now(), *record_ids],
                )
        count = cursor.rowcount
        self._audit("approve", {"count": count, "record_ids": record_ids, "approve_all": approve_all})
        return count

    def reject(self, record_ids: list[int] | None = None, reject_all: bool = False) -> int:
        """Mark pending records rejected; this still performs no filesystem operation."""
        with self._connect() as connection:
            if reject_all:
                cursor = connection.execute(
                    "UPDATE indexed_files SET status='REJECTED', updated_at=? WHERE status='PENDING'", (_now(),)
                )
            else:
                if not record_ids:
                    raise IndexerError("Provide record IDs or reject_all=True")
                placeholders = ",".join("?" for _ in record_ids)
                cursor = connection.execute(
                    f"UPDATE indexed_files SET status='REJECTED', updated_at=? WHERE status='PENDING' AND id IN ({placeholders})",
                    [_now(), *record_ids],
                )
        count = cursor.rowcount
        self._promote_duplicates()
        self._audit("reject", {"count": count, "record_ids": record_ids, "reject_all": reject_all})
        return count

    def apply(
        self,
        record_ids: list[int] | None = None,
        confirm: bool = False,
        source_root: Path | None = None,
        target_root: Path | None = None,
    ) -> dict[str, Any]:
        """Safely apply approved records after explicit confirmation."""
        if not confirm:
            raise IndexerError("Apply requires explicit confirm=True")
        statuses = {"APPROVED"}
        stored_source, stored_target = self._roots()
        if source_root is not None and source_root.expanduser().resolve() != stored_source:
            raise UnsafePathError("Apply source root does not match the indexed root")
        if target_root is not None and target_root.expanduser().resolve() != stored_target:
            raise UnsafePathError("Apply target root does not match the indexed root")
        records = self.list_queue(statuses)
        if record_ids is not None:
            wanted = set(record_ids)
            records = [record for record in records if record.id in wanted]
        applied: list[int] = []
        failed: list[dict[str, Any]] = []
        for record in records:
            try:
                self._apply_one(record)
                applied.append(record.id)
                self._promote_duplicates()
            except (OSError, IndexerError, ValueError) as exc:
                failed.append({"id": record.id, "error": type(exc).__name__})
        return {"applied": applied, "failed": failed, "count": len(applied)}

    def _apply_one(self, record: IndexRecord) -> None:
        source_root, target_root = self._roots()
        source = self._safe_child(Path(record.source_path), source_root)
        target = self._safe_child(Path(record.target_path), target_root)
        if not source.exists() or not source.is_file() or source.is_symlink():
            raise IndexerError(f"Source is unavailable or unsafe: {source}")
        target.parent.mkdir(parents=True, exist_ok=True)
        if target.exists():
            if target.is_file() and sha256_file(target) == record.sha256:
                raise IndexerError("Destination already contains the indexed file")
            target = _collision_path(target)
        temporary = target.with_name(f".{target.name}.{record.id}.tmp")
        placed = False
        try:
            shutil.copy2(source, temporary)
            if sha256_file(temporary) != record.sha256:
                raise IndexerError("Copied file hash does not match indexed hash")
            os.replace(temporary, target)
            placed = True
            source.unlink()
        except Exception as exc:
            if temporary.exists():
                temporary.unlink()
            if placed:
                self._set_status(record.id, "PARTIAL", type(exc).__name__)
            raise
        with self._connect() as connection:
            connection.execute(
                "UPDATE indexed_files SET target_path=?, status='APPLIED', updated_at=? WHERE id=?",
                (str(target), _now(), record.id),
            )
        self._audit("apply", {"id": record.id, "source_path": str(source), "target_path": str(target)})

    def _propose_target(self, source: Path, target_root: Path, modality: str, digest: str) -> Path:
        category = {"document": "documents", "image": "images", "video": "videos", "music": "music"}[modality]
        candidate = target_root / category / source.name
        return self._safe_child(candidate, target_root)

    def _hash_exists_elsewhere(self, digest: str, source_path: str) -> bool:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT 1 FROM indexed_files WHERE sha256=? AND source_path<>? "
                "AND status IN ('PENDING', 'APPROVED', 'APPLIED', 'PARTIAL') LIMIT 1",
                (digest, source_path),
            ).fetchone()
        return row is not None

    def _roots(self) -> tuple[Path, Path]:
        with self._connect() as connection:
            row = connection.execute("SELECT source_root, target_root FROM index_roots WHERE id=1").fetchone()
        if row is None:
            raise IndexerError("Indexer roots are not configured; scan before apply")
        return Path(str(row["source_root"])), Path(str(row["target_root"]))

    def _set_status(self, record_id: int, status: str, error: str | None = None) -> None:
        with self._connect() as connection:
            connection.execute(
                "UPDATE indexed_files SET status=?, error=?, updated_at=? WHERE id=?",
                (status, error, _now(), record_id),
            )

    def _promote_duplicates(self) -> None:
        with self._connect() as connection:
            hashes = connection.execute(
                "SELECT DISTINCT sha256 FROM indexed_files WHERE status='DUPLICATE' AND sha256<>''"
            ).fetchall()
            for row in hashes:
                digest = str(row["sha256"])
                owner = connection.execute(
                    "SELECT 1 FROM indexed_files WHERE sha256=? AND status IN ('PENDING','APPROVED') LIMIT 1",
                    (digest,),
                ).fetchone()
                if owner is None:
                    candidate = connection.execute(
                        "SELECT id FROM indexed_files WHERE sha256=? AND status='DUPLICATE' ORDER BY id LIMIT 1",
                        (digest,),
                    ).fetchone()
                    if candidate is not None:
                        connection.execute(
                            "UPDATE indexed_files SET status='PENDING', updated_at=? WHERE id=?",
                            (_now(), int(candidate["id"])),
                        )

    @staticmethod
    def _safe_child(path: Path, root: Path) -> Path:
        resolved = path.expanduser().resolve(strict=False)
        root_resolved = root.expanduser().resolve(strict=False)
        if resolved != root_resolved and not resolved.is_relative_to(root_resolved):
            raise UnsafePathError(f"Path escapes root: {path}")
        return resolved

    @staticmethod
    def _classify(path: Path) -> str | None:
        suffix = path.suffix.lower()
        if suffix in DOCUMENT_EXTENSIONS:
            return "document"
        if suffix in IMAGE_EXTENSIONS:
            return "image"
        if suffix in VIDEO_EXTENSIONS:
            return "video"
        if suffix in AUDIO_EXTENSIONS:
            return "music"
        return None

    def _audit(self, event: str, payload: dict[str, Any]) -> None:
        if self.audit_logger is not None:
            self.audit_logger.record(event, payload)


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    """Hash a file incrementally without loading it into memory."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def extract_metadata(
    path: Path,
    modality: str,
    analysis_hooks: Iterable[Callable[[Path, str, dict[str, Any]], dict[str, Any]]] | None = None,
) -> dict[str, Any]:
    """Extract local metadata; optional hooks are additive and never mutate files."""
    metadata: dict[str, Any] = {"filename": path.name, "suffix": path.suffix.lower()}
    if modality == "document":
        metadata = _document_metadata(path, metadata)
    elif modality == "image":
        metadata = _image_metadata(path, metadata)
    elif modality == "video":
        metadata = _ffprobe_metadata(path, metadata)
    elif modality == "music":
        metadata = _audio_metadata(path, metadata)
    for hook in analysis_hooks or []:
        try:
            additions = hook(path, modality, dict(metadata))
            if not isinstance(additions, dict):
                raise ValueError("analysis hook must return a dictionary")
            metadata.setdefault("analysis_hooks", {}).update(additions)
        except Exception as exc:
            metadata.setdefault("analysis_hook_errors", []).append(type(exc).__name__)
    return metadata


def _document_metadata(path: Path, metadata: dict[str, Any]) -> dict[str, Any]:
    if path.suffix.lower() in {".txt", ".md", ".markdown", ".rst", ".json", ".csv", ".tsv", ".html", ".htm", ".xml"}:
        try:
            with path.open("r", encoding="utf-8", errors="replace") as handle:
                preview = handle.read(4000)
            metadata["text_preview"] = preview
            metadata["text_chars"] = len(preview)
        except OSError as exc:
            metadata["metadata_error"] = type(exc).__name__
    elif path.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            metadata["page_count"] = len(reader.pages)
            metadata["text_preview"] = "\n".join((page.extract_text() or "") for page in reader.pages[:3])[:4000]
            metadata["parser"] = "pypdf"
        except ImportError:
            metadata["parser"] = "unavailable:pypdf"
        except Exception as exc:
            metadata["metadata_error"] = type(exc).__name__
    elif path.suffix.lower() == ".docx":
        try:
            from docx import Document
            document = Document(str(path))
            metadata["text_preview"] = "\n".join(paragraph.text for paragraph in document.paragraphs)[:4000]
            metadata["paragraph_count"] = len(document.paragraphs)
            metadata["parser"] = "python-docx"
        except ImportError:
            metadata["parser"] = "unavailable:python-docx"
        except Exception as exc:
            metadata["metadata_error"] = type(exc).__name__
    return metadata


def _image_metadata(path: Path, metadata: dict[str, Any]) -> dict[str, Any]:
    try:
        from PIL import Image
        with Image.open(path) as image:
            metadata.update({"width": image.width, "height": image.height, "format": image.format, "mode": image.mode})
            metadata["parser"] = "Pillow"
            exif = image.getexif()
            if exif:
                metadata["exif_keys"] = [str(key) for key in exif.keys()]
    except ImportError:
        metadata["parser"] = "unavailable:Pillow"
    except Exception as exc:
        metadata["metadata_error"] = type(exc).__name__
    return metadata


def _audio_metadata(path: Path, metadata: dict[str, Any]) -> dict[str, Any]:
    try:
        from mutagen import File
        audio = File(str(path), easy=True)
        if audio is not None:
            metadata["duration_seconds"] = getattr(audio.info, "length", None)
            for key in ("title", "artist", "album", "genre", "date"):
                if audio.get(key):
                    metadata[key] = audio.get(key)[0]
            metadata["parser"] = "mutagen"
    except ImportError:
        metadata["parser"] = "unavailable:mutagen"
    except Exception as exc:
        metadata["metadata_error"] = type(exc).__name__
    return metadata


def _ffprobe_metadata(path: Path, metadata: dict[str, Any]) -> dict[str, Any]:
    ffprobe = shutil.which("ffprobe") or shutil.which("ffprobe.exe")
    if not ffprobe:
        metadata["parser"] = "unavailable:ffprobe"
        return metadata
    try:
        result = subprocess.run(
            [ffprobe, "-v", "error", "-print_format", "json", "-show_format", "-show_streams", str(path)],
            capture_output=True, text=True, timeout=10, check=False,
        )
        if result.returncode != 0:
            metadata["metadata_error"] = "ffprobe_failed"
        else:
            payload = json.loads(result.stdout)
            metadata["duration_seconds"] = payload.get("format", {}).get("duration")
            metadata["format_name"] = payload.get("format", {}).get("format_name")
            metadata["streams"] = [
                {key: stream.get(key) for key in ("codec_type", "codec_name", "width", "height", "sample_rate", "channels")}
                for stream in payload.get("streams", [])
            ]
            metadata["parser"] = "ffprobe"
    except (OSError, subprocess.SubprocessError, json.JSONDecodeError) as exc:
        metadata["metadata_error"] = type(exc).__name__
    return metadata


def ocr_hook(language: str = "eng") -> Callable[[Path, str, dict[str, Any]], dict[str, Any]]:
    """Return an analysis hook that OCRs images into ``ocr_text``.

    Uses ``pytesseract`` + Pillow when installed; records ``ocr_parser`` as
    ``unavailable:pytesseract`` otherwise. Never mutates source files.
    """

    def hook(path: Path, modality: str, metadata: dict[str, Any]) -> dict[str, Any]:
        if modality not in ("image", "document"):
            return {}
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            return {"ocr_parser": "unavailable:pytesseract"}
        try:
            with Image.open(path) as image:
                text = pytesseract.image_to_string(image, lang=language)
        except Exception as exc:
            return {"ocr_error": type(exc).__name__, "ocr_parser": "pytesseract"}
        return {"ocr_text": text.strip() or None, "ocr_parser": "pytesseract"}

    return hook


def vision_caption_hook(
    client: Any,
    reference: str,
    prompt: str = "Describe this image in one concise sentence.",
) -> Callable[[Path, str, dict[str, Any]], dict[str, Any]]:
    """Return an analysis hook that captions images and video keyframes.

    ``client`` is any object exposing ``complete(reference, prompt, image_path=...)
    -> ModelResult`` (e.g. :class:`MultiModelClient`). Video files are reduced to
    one representative frame via ffmpeg, which is deleted after captioning.
    """

    def hook(path: Path, modality: str, metadata: dict[str, Any]) -> dict[str, Any]:
        if modality not in ("image", "video"):
            return {}
        image_path: Path = path
        frame: Path | None = None
        if modality == "video":
            frame = _extract_keyframe(path)
            if frame is None:
                return {"caption_error": "keyframe_extraction_failed"}
            image_path = frame
        try:
            result = client.complete(reference, prompt, image_path=image_path)
        finally:
            if frame is not None:
                frame.unlink(missing_ok=True)
        if not getattr(result, "ok", False):
            return {"caption_error": getattr(result, "error", None) or "caption_failed"}
        return {"caption": str(getattr(result, "text", "")).strip()}

    return hook


def _extract_keyframe(path: Path) -> Path | None:
    """Extract one representative frame to a temp file; caller must delete it."""
    ffmpeg = shutil.which("ffmpeg") or shutil.which("ffmpeg.exe")
    if not ffmpeg:
        return None
    import tempfile

    handle = tempfile.NamedTemporaryFile(suffix=".png", prefix="ais_keyframe_", delete=False)
    frame = Path(handle.name)
    handle.close()
    try:
        result = subprocess.run(
            [ffmpeg, "-y", "-ss", "0", "-i", str(path), "-vf", "thumbnail", "-frames:v", "1", str(frame)],
            capture_output=True, text=True, timeout=30, check=False,
        )
    except (OSError, subprocess.SubprocessError):
        frame.unlink(missing_ok=True)
        return None
    if result.returncode != 0 or not frame.exists() or frame.stat().st_size == 0:
        frame.unlink(missing_ok=True)
        return None
    return frame


def _collision_path(path: Path) -> Path:
    index = 2
    while True:
        candidate = path.with_name(f"{path.stem}_v{index}{path.suffix}")
        if not candidate.exists():
            return candidate
        index += 1


def _row_to_record(row: sqlite3.Row) -> IndexRecord:
    return IndexRecord(
        id=int(row["id"]), source_path=str(row["source_path"]), target_path=str(row["target_path"]),
        modality=str(row["modality"]), mime_type=row["mime_type"], size_bytes=int(row["size_bytes"]),
        modified_at=str(row["modified_at"]), sha256=str(row["sha256"]),
        metadata=json.loads(row["metadata_json"]), status=str(row["status"]), error=row["error"],
        created_at=str(row["created_at"]), updated_at=str(row["updated_at"]),
    )


def _now() -> str:
    return datetime.now(UTC).isoformat()
